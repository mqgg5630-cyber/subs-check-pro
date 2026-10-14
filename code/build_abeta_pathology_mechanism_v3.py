#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render the independent Aβ pathology/aggregation-mechanism v3 memorandum.

This builder deliberately writes a new v3 DOCX and never changes either the
historical inhibitor/neuroprotection packet or the sequence-specific v2 memo.

Usage (requires python-docx):
    python3 code/build_abeta_pathology_mechanism_v3.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "abeta_pc12_literature_review"
SOURCE = OUT / "Aβ病理机制_促聚集_AChE-PAS与PC12损伤_12条肽_独立审查_v3.md"
TARGET = OUT / "Aβ病理机制_促聚集_AChE-PAS与PC12损伤_12条肽_独立审查_v3.docx"


def inline(text: str) -> str:
    """Keep Word text readable while exposing links and removing Markdown marks."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "").replace("<br>", " ").strip()


def font(run, size: float, *, bold: bool = False, color: str = "263238") -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_cell(cell, text: str, header: bool) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(inline(text))
    font(run, 7.15 if header else 6.75, bold=header, color="FFFFFF" if header else "263238")
    if header:
        shade(cell, "0B485A")


def markdown_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if re.match(r"^\|\s*:?-{3,}", line):
            continue
        values = [part.strip() for part in line.strip().split("|")[1:-1]]
        if values:
            rows.append(values)
    return rows


def add_table(document: Document, lines: list[str]) -> None:
    rows = markdown_table(lines)
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for column in range(columns):
            set_cell(cells[column], values[column] if column < len(values) else "", row_number == 0)
        if row_number == 0:
            repeat_header(table.rows[0])
    document.add_paragraph("")


def add_paragraph(document: Document, text: str, *, style: str | None = None, size: float = 8.7, bold: bool = False, color: str = "263238") -> None:
    paragraph = document.add_paragraph(style=style) if style else document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(inline(text))
    font(run, size, bold=bold, color=color)


def build() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing source memorandum: {SOURCE}")

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)

    for style_name, size in (("Normal", 8.7), ("Title", 19), ("Heading 1", 13.5), ("Heading 2", 11.1)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
    document.styles["Normal"].paragraph_format.space_after = Pt(3)

    props = document.core_properties
    props.title = "Aβ 病理机制：促聚集、AChE–PAS 与 PC12 损伤（v3）"
    props.subject = "12 条候选短肽的机制导向文献、计算与湿实验审查"
    props.author = "Arena.ai Agent Mode"
    props.comments = "独立新建文件；不覆盖抑制/神经保护初版或促聚集 v2 文件。"

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("独立研究备忘｜Aβ 病理机制 / 促聚集 / AChE–PAS / PC12｜v3")
    font(header_run, 7.1, color="607D8B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("2026-10-14 Asia/Shanghai｜预测与对接均为待实验检验的机制假说；本文件不替代官方或原始实验记录")
    font(footer_run, 7.0, color="607D8B")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line.startswith("|"):
            block: list[str] = []
            while index < len(lines) and lines[index].startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(document, block)
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(inline(line[2:]))
            font(run, 18.5, bold=True, color="0D3443")
        elif line.startswith("## "):
            add_paragraph(document, line[3:], style="Heading 1", size=13.4, bold=True, color="07576C")
        elif line.startswith("### "):
            add_paragraph(document, line[4:], style="Heading 2", size=10.9, bold=True, color="A04431")
        elif line.strip() == "---":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run("─" * 110)
            font(run, 5.0, color="B0BEC5")
        elif line.startswith("- "):
            add_paragraph(document, line[2:], style="List Bullet", size=8.7)
        elif re.match(r"^\d+\. ", line):
            add_paragraph(document, re.sub(r"^\d+\. ", "", line), style="List Number", size=8.7)
        elif line.strip():
            add_paragraph(document, line, size=8.7)
        index += 1

    document.save(TARGET)
    print(f"Wrote {TARGET}")


if __name__ == "__main__":
    build()
