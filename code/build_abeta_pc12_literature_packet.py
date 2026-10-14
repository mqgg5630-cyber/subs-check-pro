#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Render the curated Aβ/PC12 literature-review Markdown into practical files.

The review intentionally does not claim sequence-specific conclusions until the
user's 12-peptide Word document is synced into this workspace.

Usage:
  PYTHONPATH=/tmp/accio-deps python3 code/build_abeta_pc12_literature_packet.py
"""
from __future__ import annotations

import csv
import re
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "abeta_pc12_literature_review"
MD = OUT / "Aβ聚集抑制与PC12神经保护_12条候选肽_文献与模型审查_初版.md"
DOCX = OUT / "Aβ聚集抑制与PC12神经保护_12条候选肽_文献与模型审查_初版.docx"
CSV = OUT / "核心文献索引.csv"
BIB = OUT / "references_abeta_pc12.bib"
README = OUT / "README.md"

# Deliberately compact machine-readable index; the narrative review carries
# full context and all references.
REFS = [
    ("Gautam2015", "Gautam et al.", 2015, "Strong Inhibition of Beta-Amyloid Peptide Aggregation Realized by Two-Steps Evolved Peptides", "实验/筛选/PC12", "10.1111/cbdd.12400", "25082146", "肽库→ThT/AFM→PC12 挽救；最直接的筛选先例"),
    ("Wu2023", "Wu et al.", 2023, "The interaction of peptide inhibitors and Aβ protein", "实验/对接/PC12", "10.3389/fnagi.2023.1139418", "37113572", "对接+ThT+CCK-8 PC12；适合作为小面板初筛参考"),
    ("Mallesh2023", "Mallesh et al.", 2023, "Controlling Amyloid Beta Peptide Aggregation and Toxicity by Protease-Stable Ligands", "实验/正交读出/PC12", "10.1021/acsbiomedchemau.2c00067", "37101809", "ThT/CR/CD/FTIR/TEM/PC12/apoptosis 的完整模板"),
    ("HenningKnechtel2020", "Henning-Knechtel et al.", 2020, "Designed Cell-Penetrating Peptide Inhibitors of Amyloid-beta Aggregation and Cytotoxicity", "实验/MD", "10.1016/j.xcrp.2020.100014", "", "Aβ 抑制、细胞毒性和 NMR/MD 一体化案例"),
    ("Shearman1994", "Shearman et al.", 1994, "Inhibition of PC12 cell redox activity is a specific, early indicator", "PC12 方法学", "10.1073/pnas.91.4.1470", "8108433", "MTT 是早期氧化还原读出，不能单独等同不可逆死亡"),
    ("Lemkul2012", "Lemkul & Bevan", 2012, "The Role of Molecular Simulations in the Development of Inhibitors of Amyloid β-Peptide Aggregation", "MD 综述", "10.1021/cn300091a", "", "Aβ 单体/寡聚/纤维的 docking/MD 设计与限制"),
    ("Viet2011", "Viet et al.", 2011, "Inhibition of Aggregation of Amyloid Peptides by Beta-Sheet Breaker Peptides", "MD/MM-PBSA", "10.1021/jp1116728", "21563780", "KLVFF/LPFFD 的模型机制与正控价值"),
    ("Kaur2023", "Kaur & Goyal", 2023, "Identification of new pentapeptides as potential inhibitors of amyloid–β42 aggregation", "虚拟筛选/MD", "10.1016/j.jmgm.2023.108558", "", "912 pentapeptide→docking/MM-PBSA/MD；计算候选生成参考"),
    ("Han2025", "Han et al.", 2025, "Discovery of New Pentapeptide Inhibitors Against Amyloid-β Aggregation Using Word2Vec and Molecular Simulation", "ML→MD→ThT", "10.1021/acschemneuro.4c00661", "", "最接近短肽库 Aβ 抑制剂发现的计算-实验工作流"),
    ("Grasso2020", "Grasso & Danani", 2020, "Molecular simulations of amyloid beta assemblies", "MD 综述", "10.1080/23746149.2020.1770627", "", "AA/CG/增强采样与时间尺度限制"),
    ("Familia2015", "Família et al.", 2015, "Prediction of Peptide and Protein Propensity for Amyloid Formation", "ML/APPNN", "10.1371/journal.pone.0134679", "", "候选肽自身 amyloid propensity 的传统 ANN 基线"),
    ("AmyloGram2017", "Burdukiewicz et al.", 2017, "Amyloidogenic motifs revealed by n-gram analysis", "ML/AmyloGram", "10.1038/s41598-017-13210-9", "", "6–15 aa hotspot、n-gram+RF；共识预筛"),
    ("PATH2020", "Wojciechowski & Kotulska", 2020, "PATH - Prediction of Amyloidogenicity by Threading and Machine Learning", "ML/结构 threading", "10.1038/s41598-020-64270-3", "", "hexapeptide 自身 amyloidogenicity 的补充特征"),
    ("AggBERT2023", "Perez et al.", 2023, "AggBERT: Best in Class Prediction of Hexapeptide Amyloidogenesis", "DL/ProtBERT", "10.1021/acs.jcim.3c00817", "37552230", "hexapeptide 自身聚集预测，不是 Aβ 抑制/PC12 模型"),
    ("AggNet2025", "He et al.", 2025, "AggNet: Advancing protein aggregation analysis through deep learning and protein language model", "DL/ESM2", "10.1002/pro.70031", "39840791", "amyloid/non-amyloid 和 APR；可作共识筛选"),
    ("CANYA2025", "Thompson et al.", 2025, "Massive experimental quantification allows interpretable deep learning of protein aggregation", "DL/大数据", "10.1126/sciadv.adt5111", "", "广序列空间 intrinsic aggregation；提醒模型外推限制"),
    ("PALM2026", "Eschbach et al.", 2026, "Predicting Peptide Aggregation with Protein Language Model Embeddings", "DL/ESM2", "10.1021/acs.jcim.6c00577", "42049499", "推荐优先做自身聚集/APR 预筛；非 Aβ 抑制器"),
    ("TANGO2004", "Fernandez-Escamilla et al.", 2004, "Prediction of sequence-dependent and mutational effects on aggregation", "物理/基线", "10.1038/nbt1012", "", "β 聚集风险 sanity check"),
    ("PASTA2014", "Walsh et al.", 2014, "PASTA 2.0: an improved server for protein aggregation prediction", "物理/基线", "10.1093/nar/gku399", "", "β-pairing/aggregation-prone 区域基线"),
]


def clean_inline(text: str) -> str:
    """Keep the Word output readable while retaining both link label and URL."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("<br>", "")
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def set_font(run, name: str = "Microsoft YaHei", size: float | None = None, bold: bool | None = None, color: str | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, value: str, *, header: bool = False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(clean_inline(value))
    set_font(run, size=8.2 if header else 7.7, bold=header, color="FFFFFF" if header else "263238")
    if header:
        shade(cell, "0F3445")


def add_table(doc: Document, raw_lines: list[str]):
    rows = []
    for line in raw_lines:
        if re.match(r"^\|\s*:?-{3,}", line):
            continue
        parts = [p.strip() for p in line.strip().split("|")[1:-1]]
        if parts:
            rows.append(parts)
    if not rows:
        return
    cols = max(map(len, rows))
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx in range(cols):
            set_cell_text(cells[cidx], row[cidx] if cidx < len(row) else "", header=(ridx == 0))
    doc.add_paragraph("")


def add_markdown_document() -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.56)
    section.right_margin = Inches(0.56)
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.52)

    styles = doc.styles
    for style_name, size in [("Normal", 9.2), ("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 10.5)]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
    styles["Normal"].paragraph_format.space_after = Pt(3)

    # Header and footer
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = header.add_run("Aβ / PC12 文献与模型审查｜初版｜仅供研究设计参考")
    set_font(hrun, size=7.5, color="607D8B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = footer.add_run("生成时间：2026-10-14 Asia/Shanghai｜收到 12 条肽原始序列与实验方案后需做特异性复核")
    set_font(frun, size=7.5, color="607D8B")

    lines = MD.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        # Markdown table block
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(clean_inline(line[2:]))
            set_font(run, size=21, bold=True, color="102A3A")
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            run = p.add_run(clean_inline(line[3:]))
            set_font(run, size=15, bold=True, color="0F6070")
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            run = p.add_run(clean_inline(line[4:]))
            set_font(run, size=12, bold=True, color="A04431")
        elif line.strip() == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run("─" * 120)
            for run in p.runs:
                set_font(run, size=6.5, color="B0BEC5")
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(clean_inline(line[2:]))
            set_font(run, size=9.2)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(clean_inline(re.sub(r"^\d+\. ", "", line)))
            set_font(run, size=9.2)
        elif line.strip():
            p = doc.add_paragraph()
            run = p.add_run(clean_inline(line))
            set_font(run, size=9.2)
        i += 1
    doc.save(DOCX)


def write_index_and_bib() -> None:
    with CSV.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(["cite_key", "first_author", "year", "title", "category", "doi", "doi_url", "pmid", "why_it_matters"])
        for key, author, year, title, category, doi, pmid, note in REFS:
            writer.writerow([key, author, year, title, category, doi, f"https://doi.org/{doi}", pmid, note])
    bib = []
    for key, author, year, title, category, doi, pmid, note in REFS:
        fields = [
            f"@misc{{{key},",
            f"  author = {{{author}}},",
            f"  title = {{{title}}},",
            f"  year = {{{year}}},",
            f"  doi = {{{doi}}},",
            f"  url = {{https://doi.org/{doi}}},",
            f"  note = {{{category}; {note}}},",
        ]
        if pmid:
            fields.append(f"  pmid = {{{pmid}}},")
        fields.append("}")
        bib.append("\n".join(fields))
    BIB.write_text("\n\n".join(bib) + "\n", encoding="utf-8")


def write_readme(timestamp: str) -> None:
    README.write_text(
        f"""# Aβ 聚集抑制与 PC12 神经保护｜文献与模型审查（初版）

**生成时间：{timestamp}**

## 文件

- `Aβ聚集抑制与PC12神经保护_12条候选肽_文献与模型审查_初版.md` — 完整、可搜索的中文审查正文。
- `Aβ聚集抑制与PC12神经保护_12条候选肽_文献与模型审查_初版.docx` — 便于转发、批注和放入课题资料的 Word 版本。
- `核心文献索引.csv` — 19 条核心文献的 DOI/PMID/用途索引，可导入表格软件。
- `references_abeta_pc12.bib` — 用于 Zotero/BibTeX 的简化引用索引；正式投稿前应从 DOI 元数据服务核对所有作者、卷期页码。
- `12条候选肽_序列与实验映射模板.csv` — 待拿到 12 条序列后填充的映射模板。

## 当前边界

用户给出的 Windows `.docx` 路径未同步到此工作区，因此本版不能把任何结论说成针对具体 12 条肽。请上传该文档或填写序列映射 CSV；收到后即可进行逐条序列-文献-模型-实验条件映射。

## 重建

```bash
PYTHONPATH=/tmp/accio-deps python3 code/build_abeta_pc12_literature_packet.py
```

脚本需要 `python-docx`；Markdown 与 CSV 不依赖该库。
""",
        encoding="utf-8",
    )


def main() -> None:
    if not MD.exists():
        raise FileNotFoundError(MD)
    timestamp = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M Asia/Shanghai")
    add_markdown_document()
    write_index_and_bib()
    write_readme(timestamp)
    for path in (DOCX, CSV, BIB, README):
        print(f"built {path.relative_to(ROOT)} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
