#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the independent, sequence-specific Aβ/AChE-PAS/PC12 addendum.

This renderer deliberately treats the supplied docking / prediction results as
hypotheses.  It does not alter the earlier inhibitor/neuroprotection packet.

Usage:
  PYTHONPATH=/tmp/accio-deps python3 code/build_abeta_pas_12_peptide_addendum.py
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "abeta_pc12_literature_review"
MD = OUT / "Aβ促聚集_AChE-PAS_12条对接肽_序列特异性复核与湿实验计划_v2.md"
DOCX = OUT / "Aβ促聚集_AChE-PAS_12条对接肽_序列特异性复核与湿实验计划_v2.docx"
MATRIX = OUT / "12条AChE对接肽_序列特异性复核矩阵_v2.csv"
REFS = OUT / "AChE_PAS_Aβ_PC12_关键文献索引_v2.csv"
README = OUT / "README_12条对接肽加注_v2.md"

# The table is deliberately small and traceable.  It complements—not replaces—
# the numbered reference list in the narrative memorandum.
REFERENCE_ROWS = [
    (
        "Inestrosa1996",
        "Inestrosa NC et al.",
        "1996",
        "Acetylcholinesterase Accelerates Assembly of Amyloid-β-Peptides into Alzheimer's Fibrils: Possible Role of the Peripheral Site of the Enzyme",
        "Neuron 16:881–891",
        "10.1016/S0896-6273(00)80108-7",
        "AChE can enhance Aβ fibrillogenesis; propidium attenuated the effect whereas the active-site inhibitor edrophonium did not. This is the central directionality control.",
    ),
    (
        "DeFerrari2001",
        "De Ferrari GV et al.",
        "2001",
        "A structural motif of acetylcholinesterase that promotes amyloid beta-peptide fibril formation",
        "Biochemistry 40:10447–10457",
        "10.1021/bi0101392",
        "A hydrophobic AChE motif near PAS can influence Aβ fibrillation; it does not establish that arbitrary short PAS-binding peptides promote aggregation.",
    ),
    (
        "Alvarez1998",
        "Alvarez A et al.",
        "1998",
        "Stable complexes involving acetylcholinesterase and amyloid-beta peptide change the biochemical properties of the enzyme and increase the neurotoxicity of Alzheimer's fibrils",
        "J Neurosci 18:3213–3223",
        "",
        "Supports treating AChE–Aβ aggregate toxicity as a distinct experimental condition from Aβ alone.",
    ),
    (
        "MunozInestrosa1999",
        "Muñoz FJ; Inestrosa NC",
        "1999",
        "Neurotoxicity of acetylcholinesterase amyloid beta-peptide aggregates is dependent on the type of Aβ peptide and the AChE concentration present in the complexes",
        "FEBS Lett",
        "10.1016/S0014-5793(99)00468-8",
        "Supports separate Aβ, AChE–Aβ, and AChE–Aβ–peptide conditions in PC12 assays.",
    ),
    (
        "PDB4EY6",
        "RCSB Protein Data Bank",
        "accessed 2026-10-14",
        "Crystal Structure of Recombinant Human Acetylcholinesterase in Complex with (-)-galantamine (PDB 4EY6)",
        "RCSB PDB entry 4EY6",
        "",
        "Confirms that 4EY6 is a human AChE galantamine co-crystal; ligand/heteroatom handling must be disclosed.",
    ),
    (
        "Rathore2025",
        "Rathore AS et al.",
        "2025",
        "A large language model for predicting neurotoxic peptides and neurotoxins",
        "Protein Science 34:e70200",
        "10.1002/pro.70200",
        "NTxPred2 classifies generic neurotoxic peptide labels; it is not a PC12 viability, Aβ-coexposure, or dose-response model.",
    ),
    (
        "Thompson2025",
        "Thompson M et al.",
        "2025",
        "Massive experimental quantification allows interpretable deep learning of protein aggregation",
        "Science Advances 11:eadt5111",
        "10.1126/sciadv.adt5111",
        "CANYA was trained from random 20-aa sequences in a Sup35N-fusion selection; short 7–9-aa free peptides are a material domain shift.",
    ),
    (
        "Perez2023",
        "Perez R et al.",
        "2023",
        "AggBERT: Best in Class Prediction of Hexapeptide Amyloidogenesis with a Semi-Supervised ProtBERT Model",
        "J Chem Inf Model 63:5727–5733",
        "10.1021/acs.jcim.3c00817",
        "A hexapeptide intrinsic-amyloidogenicity model; at most a sliding-window self-aggregation annotation for this project.",
    ),
    (
        "Bansal2025",
        "Bansal A et al.",
        "2025",
        "Machine Learning–Based Bioactivity Prediction and Descriptor-Guided Rational Design of Amyloid-β Aggregation Inhibitors",
        "ACS Chemical Neuroscience 16:4351–4363",
        "10.1021/acschemneuro.5c00649",
        "Small-molecule Aβ inhibitor/IC50 model; wrong molecule class and opposite biological direction for the present peptides.",
    ),
    (
        "Shearman1994",
        "Shearman MS et al.",
        "1994",
        "Inhibition of PC12 cell redox activity is a specific, early indicator of the mechanism of beta-amyloid-mediated cell death",
        "PNAS 91:1470–1474",
        "10.1073/pnas.91.4.1470",
        "MTT/CCK-8-type redox viability readouts are useful but insufficient alone to claim irreversible cell death.",
    ),
]


def clean_inline(text: str) -> str:
    """Render Markdown links as visible label + URL, avoiding hidden claims."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("<br>", " ")
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def set_font(run, *, size: float, bold: bool = False, color: str = "263238") -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell(cell, text: str, *, header: bool = False) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(clean_inline(text))
    set_font(run, size=7.2 if header else 6.9, bold=header, color="FFFFFF" if header else "263238")
    if header:
        shade(cell, "0B485A")


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if re.match(r"^\|\s*:?-{3,}", line):
            continue
        cells = [part.strip() for part in line.strip().split("|")[1:-1]]
        if cells:
            rows.append(cells)
    return rows


def add_table(doc: Document, raw_lines: list[str]) -> None:
    rows = parse_table(raw_lines)
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index in range(col_count):
            set_cell(cells[col_index], values[col_index] if col_index < len(values) else "", header=(row_index == 0))
        if row_index == 0:
            set_repeat_table_header(table.rows[0])
    doc.add_paragraph("")


def add_hr(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("─" * 120)
    set_font(run, size=5.5, color="B0BEC5")


def render_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = Inches(0.52)
    section.right_margin = Inches(0.52)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)

    for style_name, size in (("Normal", 8.8), ("Title", 20), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
    doc.styles["Normal"].paragraph_format.space_after = Pt(3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = header.add_run("独立研究备忘｜Aβ 促聚集 / AChE–PAS / PC12｜序列特异性复核 v2")
    set_font(hrun, size=7.2, color="607D8B")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = footer.add_run("2026-10-14 Asia/Shanghai｜计算预测与 docking/MD 均为待实验检验的假说，不作为 SCI 主文既成结论")
    set_font(frun, size=7.2, color="607D8B")

    lines = MD.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(clean_inline(line[2:]))
            set_font(run, size=19, bold=True, color="0D3443")
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.space_before = Pt(7)
            run = p.add_run(clean_inline(line[3:]))
            set_font(run, size=13.5, bold=True, color="07576C")
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.space_before = Pt(4)
            run = p.add_run(clean_inline(line[4:]))
            set_font(run, size=11, bold=True, color="A04431")
        elif line.strip() == "---":
            add_hr(doc)
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1.5)
            run = p.add_run(clean_inline(line[2:]))
            set_font(run, size=8.8)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(1.5)
            run = p.add_run(clean_inline(re.sub(r"^\d+\. ", "", line)))
            set_font(run, size=8.8)
        elif line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_inline(line))
            set_font(run, size=8.8)
        i += 1
    doc.save(DOCX)


def write_references() -> None:
    with REFS.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(["cite_key", "authors", "year", "title", "venue", "doi", "doi_url", "why_used_in_v2"])
        for row in REFERENCE_ROWS:
            cite_key, authors, year, title, venue, doi, reason = row
            writer.writerow([cite_key, authors, year, title, venue, doi, f"https://doi.org/{doi}" if doi else "", reason])


def write_readme() -> None:
    README.write_text(
        """# 12 条 AChE 对接肽｜促 Aβ 聚集 / PC12 神经毒性序列特异性加注（v2）

**定位：独立研究备忘，不属于 SCI 主文。**

本目录原有的“初版”文献包以 Aβ 聚集抑制 / 神经保护为问题背景。此次新增的 v2 不替换初版，而是依据用户粘贴的 12 条肽、对接、MD 和拟议实验，专门复核相反的假说：外源短肽是否会经 AChE（尤其 PAS）影响 Aβ 聚集与 PC12 损伤。

## 新增文件

- `Aβ促聚集_AChE-PAS_12条对接肽_序列特异性复核与湿实验计划_v2.md` — 可搜索的完整中文备忘；含逻辑校正、12 条逐条分层、模型适用性、可证伪实验和文献。
- `Aβ促聚集_AChE-PAS_12条对接肽_序列特异性复核与湿实验计划_v2.docx` — 用于转发、批注的 Word 版本。
- `12条AChE对接肽_序列特异性复核矩阵_v2.csv` — 12 条序列的结构接触、Cys/自聚集旗标、角色和判读限制；含的粗略净电荷是“游离 N 端 / C 端酰胺”假设下的近似值，非实验测量。
- `AChE_PAS_Aβ_PC12_关键文献索引_v2.csv` — 对本 v2 关键机制判断的可审计文献索引。

## 核心修正

1. PAS docking 不可直接推导为 AChE 介导的促 Aβ 聚集；它也可能竞争 / 阻断 AChE–Aβ。
2. NTxPred2 阳性不等价于 PC12 神经毒性；CANYA/PALM/AggBERT 等自身聚集工具不等价于促 Aβ。
3. 4EY6 是 galantamine 共晶；受体预处理、末端酰胺和 Cys 氧化态必须与真实合成肽一致。
4. FLLHTTR、YLSLLQR、ALLLHRC 是“PAS 主候选 1/2 + 峡内/催化区相关比较肽”的实用三肽配置；ALLLHRC 同时触及 Tyr124 和 Ser203，因而不是严格的 PAS 阴性或催化位点专一对照。
5. 只有配对的 Aβ、AChE、PAS 阻断、肽-alone 和正交聚集物/细胞控制才能区分促进、竞争、ThT 干扰和直接毒性。

## 重建

```bash
PYTHONPATH=/tmp/accio-deps python3 code/build_abeta_pas_12_peptide_addendum.py
```

脚本需要 `python-docx`。它只生成 v2 的 DOCX、参考索引和 README；不修改初版文献包。
""",
        encoding="utf-8",
    )


def main() -> None:
    if not MD.exists():
        raise FileNotFoundError(MD)
    if not MATRIX.exists():
        raise FileNotFoundError(MATRIX)
    render_docx()
    write_references()
    write_readme()
    for path in (DOCX, MATRIX, REFS, README):
        print(f"built {path.relative_to(ROOT)} ({path.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
